"""Read and write file contents (text, PDF, Excel, Word)."""
from __future__ import annotations

import asyncio
from pathlib import Path
from typing import Any

from ...config import settings
from ...security import PathError, safe_resolve
from ..base import Skill


def _root() -> Path:
    return Path(settings.file_root).expanduser().resolve()


def _resolve(path: str) -> Path:
    return safe_resolve(path, _root())


# ─── read ──────────────────────────────────────────────────────────────────────

class ReadFileSkill(Skill):
    name = "read_file"
    description = "Read the contents of a file (text, PDF, Excel, Word)."
    tool_description = (
        "Read the contents of any file: plain text, Python, CSV, JSON, PDF, Excel (.xlsx), "
        "Word (.docx). Returns the extracted text."
    )
    # "open" intentionally excluded — that's handled by open_app.
    patterns = [
        r"^\s*(?:read|show|cat)\s+(?:file\s+)?(?P<path>.+?)\s*[!.\?]*\s*$",
    ]
    args_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "Path to the file."},
            "max_chars": {"type": "integer", "default": 12000},
        },
        "required": ["path"],
    }

    async def run(self, text: str, match: Any | None = None) -> dict[str, Any]:
        path = match.group("path").strip() if match else text.strip()
        return await self.run_tool({"path": path})

    async def run_tool(self, args: dict[str, Any]) -> dict[str, Any]:
        try:
            target = _resolve(args.get("path") or "")
        except PathError as e:
            return {"ok": False, "reply": str(e)}
        if not target.exists():
            return {"ok": False, "reply": f"File not found: {target}"}
        if not target.is_file():
            return {"ok": False, "reply": f"{target} is not a file."}

        max_chars = int(args.get("max_chars") or 12000)
        ext = target.suffix.lower()

        try:
            content = await asyncio.to_thread(_extract, target, ext, max_chars)
        except Exception as e:
            return {"ok": False, "reply": f"Could not read {target.name}: {e}"}

        truncated = len(content) >= max_chars
        word_count = len(content.split())
        reply = f"Read {target.name} ({word_count} words)."
        if truncated:
            reply += " (truncated to fit context)"
        return {
            "ok": True,
            "reply": reply,
            "path": str(target),
            "content": content,
            "truncated": truncated,
        }


def _extract(path: Path, ext: str, max_chars: int) -> str:
    if ext == ".pdf":
        return _read_pdf(path, max_chars)
    if ext in (".xlsx", ".xls", ".xlsm"):
        return _read_excel(path, max_chars)
    if ext in (".docx",):
        return _read_docx(path, max_chars)
    # plain text fallback (txt, py, js, json, md, csv, yaml, toml, …)
    text = path.read_text(encoding="utf-8", errors="replace")
    return text[:max_chars]


def _read_pdf(path: Path, max_chars: int) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(str(path))
    parts: list[str] = []
    for page in doc:
        parts.append(page.get_text())
        if sum(len(p) for p in parts) >= max_chars:
            break
    doc.close()
    return "\n".join(parts)[:max_chars]


def _read_excel(path: Path, max_chars: int) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    rows: list[str] = []
    for sheet in wb.worksheets:
        rows.append(f"=== {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            rows.append("\t".join("" if v is None else str(v) for v in row))
            if sum(len(r) for r in rows) >= max_chars:
                return "\n".join(rows)[:max_chars]
    wb.close()
    return "\n".join(rows)[:max_chars]


def _read_docx(path: Path, max_chars: int) -> str:
    from docx import Document
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    return text[:max_chars]


# ─── write ─────────────────────────────────────────────────────────────────────

class WriteFileSkill(Skill):
    name = "write_file"
    description = "Overwrite or patch a text file."
    tool_description = (
        "Write (or overwrite) a plain-text file with new content. "
        "For partial edits, provide old_text and new_text to do a find-and-replace patch. "
        "Only plain-text formats (txt, py, md, json, csv, etc.) are supported."
    )
    requires_confirmation = True
    patterns: list[str] = []  # LLM-only — too risky for regex
    args_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "content": {
                "type": "string",
                "description": "Full new content. Used when overwriting the whole file.",
            },
            "old_text": {
                "type": "string",
                "description": "Exact text to find (for patch mode).",
            },
            "new_text": {
                "type": "string",
                "description": "Replacement text (for patch mode).",
            },
            "confirm": {
                "type": "boolean",
                "description": "Must be true to actually write.",
            },
        },
        "required": ["path", "confirm"],
    }

    async def run(self, text: str, match: Any | None = None) -> dict[str, Any]:
        return {"ok": False, "reply": "File writes only via tool with explicit confirm=true."}

    async def run_tool(self, args: dict[str, Any]) -> dict[str, Any]:
        if not (args.get("confirm") or settings.allow_destructive):
            return {"ok": False, "reply": "Refusing to write without confirm=true.", "needs_confirm": True}

        try:
            target = _resolve(args.get("path") or "")
        except PathError as e:
            return {"ok": False, "reply": str(e)}

        ext = target.suffix.lower()
        if ext in (".pdf", ".xlsx", ".xls", ".docx"):
            return {"ok": False, "reply": f"Cannot write binary format {ext} — only plain text files."}

        try:
            if args.get("old_text") is not None:
                # patch mode
                existing = await asyncio.to_thread(
                    target.read_text, encoding="utf-8", errors="replace"
                )
                if args["old_text"] not in existing:
                    return {"ok": False, "reply": "old_text not found in file — no changes made."}
                patched = existing.replace(args["old_text"], args.get("new_text") or "", 1)
                await asyncio.to_thread(
                    target.write_text, patched, encoding="utf-8"
                )
                return {"ok": True, "reply": f"Patched {target.name}.", "path": str(target)}
            else:
                # full overwrite
                content = args.get("content") or ""
                target.parent.mkdir(parents=True, exist_ok=True)
                await asyncio.to_thread(target.write_text, content, encoding="utf-8")
                return {
                    "ok": True,
                    "reply": f"Wrote {target.name} ({len(content.split())} words).",
                    "path": str(target),
                }
        except OSError as e:
            return {"ok": False, "reply": f"Write failed: {e}"}
